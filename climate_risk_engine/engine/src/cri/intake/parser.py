"""Parse a filled-in client intake Excel file into Company objects.

Supports two formats:
  1. CRI Intake Template — workbook with "Company" and "Assets" sheets
  2. Financial Summary   — any single sheet with name/sector/revenue columns
     (the engine synthesises a headquarters asset so physical risk can run)
"""

import re
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from cri.data.schemas import (
    Asset,
    Commodity,
    Company,
    EmissionsProfile,
    Financials,
)

# ---------------------------------------------------------------------------
# Financial-format column aliases
# ---------------------------------------------------------------------------
_COL = {
    "company_name": ["company_name", "company", "name", "issuer", "borrower", "entity", "firm", "client"],
    "sector":       ["sector", "industry", "gics_sector", "sector_category", "business", "sub_sector"],
    "hq_region":    ["hq_region", "region", "country", "hq_country", "geography", "location", "domicile", "headquarters"],
    "revenue_musd": ["revenue_musd", "revenue", "revm", "rev_m", "rev", "sales", "turnover", "net_revenue", "total_revenue", "revenues_musd"],
    "ebitda_musd":  ["ebitda_musd", "ebitda", "ebitdam", "ebitda_m"],
    "market_cap":   ["market_cap", "mktcap", "mktcapm", "mcap", "market_capitalization", "equity_value", "market_value"],
    "net_debt":     ["net_debt", "netdebt", "netdebtm", "net_debt_musd", "debt_net", "net_financial_debt"],
    "wacc_base":    ["wacc", "waccpct", "wacc_pct", "discount_rate", "cost_of_capital", "hurdle_rate"],
    "latitude":     ["latitude", "lat", "asset_lat", "hq_lat"],
    "longitude":    ["longitude", "lon", "lng", "asset_lon", "hq_lon"],
}

# Sector string → closest Commodity enum value for synthetic HQ asset
_SECTOR_COMMODITY = {
    "beverages": Commodity.BEVERAGES,
    "beer": Commodity.BEVERAGES,
    "brewing": Commodity.BEVERAGES,
    "drinks": Commodity.BEVERAGES,
    "food": Commodity.FOOD,
    "food_beverage": Commodity.BEVERAGES,
    "food & beverage": Commodity.BEVERAGES,
    "fmcg": Commodity.FOOD,
    "agriculture": Commodity.AGRICULTURE,
    "farming": Commodity.AGRICULTURE,
    "chemicals": Commodity.CHEMICALS,
    "specialty chemicals": Commodity.CHEMICALS,
    "oil": Commodity.CRUDE_OIL,
    "oil_gas": Commodity.CRUDE_OIL,
    "oil & gas": Commodity.CRUDE_OIL,
    "petroleum": Commodity.CRUDE_OIL,
    "natural gas": Commodity.NATURAL_GAS,
    "gas": Commodity.NATURAL_GAS,
    "coal": Commodity.COAL_THERMAL,
    "utilities": Commodity.ELECTRICITY,
    "electric": Commodity.ELECTRICITY,
    "power": Commodity.ELECTRICITY,
    "energy": Commodity.ELECTRICITY,
    "real estate": Commodity.REAL_ESTATE,
    "property": Commodity.REAL_ESTATE,
    "reit": Commodity.REAL_ESTATE,
    "cement": Commodity.CEMENT,
    "construction": Commodity.CEMENT,
    "mining": Commodity.IRON_ORE,
    "metals": Commodity.IRON_ORE,
    "metals & mining": Commodity.IRON_ORE,
    "aluminium": Commodity.ALUMINIUM,
    "aluminum": Commodity.ALUMINIUM,
    "copper": Commodity.COPPER,
    "financials": Commodity.FINANCIAL_SERVICES,
    "banking": Commodity.FINANCIAL_SERVICES,
    "insurance": Commodity.FINANCIAL_SERVICES,
    "financial services": Commodity.FINANCIAL_SERVICES,
    "retail": Commodity.RETAIL,
    "consumer": Commodity.RETAIL,
}

# Region code → (lat, lon) centroid for synthetic HQ asset
_REGION_CENTROIDS: dict[str, tuple[float, float]] = {
    "global": (20.0, 0.0),
    "us": (37.1, -95.7),   "usa": (37.1, -95.7),
    "us-ny": (40.7, -74.0), "us-ca": (37.8, -122.4), "us-tx": (30.3, -97.7),
    "us-fl": (27.6, -81.5), "us-ga": (33.4, -84.4), "us-mn": (46.4, -93.3),
    "gb": (52.5, -1.9),    "gb-eng": (52.5, -1.9), "gb-wls": (52.1, -3.8),
    "de": (51.2, 10.5),    "de-hh": (53.6, 10.0),
    "fr": (46.2, 2.2),     "nl": (52.1, 5.3),    "nl-nh": (52.3, 4.9),
    "se": (60.1, 18.6),    "no": (60.5, 8.5),    "dk": (56.3, 9.5),
    "fi": (61.9, 25.7),    "es": (40.4, -3.7),   "es-ct": (41.6, 1.5),
    "it": (41.9, 12.6),    "pt": (39.4, -8.2),   "ch": (46.8, 8.2),
    "au": (-25.3, 133.8),  "au-wa": (-25.0, 122.0), "au-sa": (-30.0, 135.8),
    "au-qld": (-20.9, 142.7), "au-nsw": (-32.1, 146.9), "au-vic": (-36.9, 144.7),
    "nz": (-40.9, 174.9),  "ca": (56.1, -106.3), "ca-ab": (53.9, -116.6),
    "ca-qc": (52.9, -73.5), "jp": (36.2, 138.3), "cn": (35.9, 104.2),
    "in": (20.6, 78.9),    "in-ap": (15.9, 79.7), "in-gj": (22.3, 71.2),
    "in-ka": (15.3, 75.7), "in-mp": (22.9, 78.7), "in-od": (20.9, 84.2),
    "in-rj": (27.0, 74.2), "br": (-10.8, -53.0), "za": (-29.0, 25.1),
    "ng": (9.1, 8.7),      "sa": (23.9, 45.1),   "ae": (23.4, 53.8),
    "sg": (1.35, 103.8),   "kr": (35.9, 127.8),
    "mn-01": (47.9, 106.9),
}


def _slug(text: str) -> str:
    """Lowercase, replace non-alnum with hyphen."""
    return re.sub(r'[^a-z0-9]+', '-', str(text).lower()).strip('-')


def _normalise_col(raw: str) -> str:
    """Normalise a column header to lowercase-no-spaces for alias matching."""
    return re.sub(r'[\s_\-]+', '_', raw.strip().lower())


def _find_col(headers: list[str], canonical: str) -> Optional[str]:
    """Return the actual column header that maps to canonical, or None."""
    aliases = [_normalise_col(a) for a in _COL.get(canonical, [])]
    for h in headers:
        if _normalise_col(h) in aliases:
            return h
    return None


def _sector_to_commodity(sector: str) -> Commodity:
    s = sector.strip().lower()
    for key, commodity in _SECTOR_COMMODITY.items():
        if key in s:
            return commodity
    return Commodity.MANUFACTURING


def _region_centroid(region: str) -> tuple[float, float]:
    key = region.strip().lower().replace(' ', '-')
    return _REGION_CENTROIDS.get(key, _REGION_CENTROIDS["global"])


def _safe_float(v, default: float = 0.0) -> float:
    try:
        return float(v) if v is not None else default
    except (TypeError, ValueError):
        return default


def _detect_financial_sheet(wb) -> Optional[str]:
    """
    Return the name of a sheet that looks like a financial summary table,
    or None if no such sheet is found.
    A sheet qualifies if its first non-empty row has at least one column
    that aliases to 'company_name' AND at least one that aliases to 'sector'.
    """
    name_aliases  = {_normalise_col(a) for a in _COL["company_name"]}
    sector_aliases = {_normalise_col(a) for a in _COL["sector"]}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        # Find first non-empty row (headers)
        for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
            raw_headers = [str(c).strip() for c in row if c is not None]
            if not raw_headers:
                continue
            norm = {_normalise_col(h) for h in raw_headers}
            if norm & name_aliases and norm & sector_aliases:
                return sheet_name
            break  # only check first non-empty header row per sheet

    return None


# Mapping of common aliases to canonical commodity names
COMMODITY_ALIASES = {
    "iron ore": "iron_ore",
    "iron_ore": "iron_ore",
    "copper": "copper",
    "aluminium": "aluminium",
    "aluminum": "aluminium",
    "coal thermal": "coal_thermal",
    "coal_thermal": "coal_thermal",
    "thermal coal": "coal_thermal",
    "coal metallurgical": "coal_metallurgical",
    "coal_metallurgical": "coal_metallurgical",
    "metallurgical coal": "coal_metallurgical",
    "crude oil": "crude_oil",
    "crude_oil": "crude_oil",
    "natural gas": "natural_gas",
    "natural_gas": "natural_gas",
    "refined products": "refined_products",
    "refined_products": "refined_products",
    "cement": "cement",
    "electricity": "electricity",
}


def _normalize_commodity(value: Optional[str]) -> str:
    """
    Normalize commodity string to canonical form.
    Raises ValueError if not recognized.
    """
    if not value:
        raise ValueError("Commodity cannot be empty")

    normalized = value.strip().lower()
    if normalized in COMMODITY_ALIASES:
        return COMMODITY_ALIASES[normalized]

    raise ValueError(
        f"Unknown commodity: {value!r}. Valid values: "
        f"{', '.join(sorted(set(COMMODITY_ALIASES.values())))}"
    )


def _read_sheet(wb, sheet_name: str):
    """Read worksheet and return list of dicts (excluding header row)."""
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet {sheet_name!r} not found in workbook")

    ws = wb[sheet_name]

    # Read headers from first row
    headers = []
    for cell in ws[1]:
        if cell.value:
            headers.append(str(cell.value).strip())

    if not headers:
        raise ValueError(f"Sheet {sheet_name!r} has no headers")

    # Read data rows
    rows = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
        # Skip completely empty rows
        if all(cell.value is None for cell in row):
            continue

        row_data = {}
        for col_idx, cell in enumerate(row):
            if col_idx < len(headers):
                # Get the value, handling different types
                value = cell.value
                row_data[headers[col_idx]] = value

        rows.append((row_idx, row_data))

    return headers, rows


def parse_excel(path: str | Path) -> list[Company]:
    """
    Parse a client intake Excel workbook and return a list of Company objects.

    Raises ValueError with clear row/column information on validation errors.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    wb = load_workbook(path, data_only=True)

    # Read Company sheet
    company_headers, company_rows = _read_sheet(wb, "Company")

    # Read Assets sheet
    asset_headers, asset_rows = _read_sheet(wb, "Assets")

    # Parse companies
    companies_by_id = {}

    for row_idx, company_data in company_rows:
        try:
            company_id_raw = company_data.get("company_id")
            if company_id_raw is None:
                raise ValueError("company_id is required")
            company_id = str(company_id_raw).strip()
            if not company_id:
                raise ValueError("company_id is required")

            company_name = str(company_data.get("company_name", "")).strip()
            if not company_name:
                raise ValueError("company_name is required")

            sector = str(company_data.get("sector", "")).strip()
            if not sector:
                raise ValueError("sector is required")

            hq_region = str(company_data.get("hq_region", "global")).strip() or "global"

            # Parse financial data
            total_debt = float(company_data.get("total_debt_musd") or 0.0)
            cash = float(company_data.get("cash_musd") or 0.0)
            shares_outstanding = float(company_data.get("shares_outstanding_m") or 1.0)
            share_price = float(company_data.get("current_share_price") or 0.0)
            revenue = float(company_data.get("revenue_musd") or 0.0)
            ebitda = float(company_data.get("ebitda_musd") or 0.0)

            # Calculate net_debt and market_cap
            net_debt = total_debt - cash
            market_cap = shares_outstanding * share_price if share_price > 0 else None

            financials = Financials(
                revenue=revenue,
                ebitda=ebitda,
                capex=0.0,  # Not in template
                net_debt=net_debt,
                shares_outstanding=shares_outstanding,
                market_cap=market_cap,
            )

            company = Company(
                id=company_id,
                name=company_name,
                sector=sector,
                hq_region=hq_region,
                financials=financials,
            )
            companies_by_id[company_id] = company

        except (ValueError, TypeError) as e:
            raise ValueError(f"Company sheet row {row_idx}: {e}")

    # Parse assets and attach to companies
    for row_idx, asset_data in asset_rows:
        try:
            company_id = str(asset_data.get("company_id", "")).strip()
            if not company_id:
                raise ValueError("company_id is required")

            if company_id not in companies_by_id:
                raise ValueError(f"company_id {company_id!r} not found in Company sheet")

            asset_id = str(asset_data.get("asset_id", "")).strip()
            if not asset_id:
                raise ValueError("asset_id is required")

            asset_name = str(asset_data.get("asset_name", "")).strip()
            if not asset_name:
                raise ValueError("asset_name is required")

            # Normalize commodity
            commodity_str = asset_data.get("commodity", "")
            commodity_str = _normalize_commodity(commodity_str)
            commodity = Commodity(commodity_str)

            region = str(asset_data.get("region", "global")).strip() or "global"

            latitude = float(asset_data.get("latitude") or 0.0)
            longitude = float(asset_data.get("longitude") or 0.0)

            baseline_production = float(asset_data.get("baseline_production") or 0.0)
            production_unit = str(asset_data.get("production_unit", "tonnes")).strip() or "tonnes"
            baseline_unit_cost = float(asset_data.get("baseline_unit_cost") or 0.0)
            energy_cost_share = float(asset_data.get("energy_cost_share") or 0.3)
            carrying_value = float(asset_data.get("carrying_value_musd") or 0.0)
            remaining_life = asset_data.get("remaining_life_years")
            if remaining_life is not None:
                remaining_life = int(remaining_life)

            scope1 = float(asset_data.get("scope1_intensity") or 0.0)
            scope2 = float(asset_data.get("scope2_intensity") or 0.0)
            scope3 = float(asset_data.get("scope3_intensity") or 0.0)
            carbon_coverage = float(asset_data.get("carbon_price_coverage") or 1.0)
            free_alloc = float(asset_data.get("free_allocation") or 0.0)

            emissions = EmissionsProfile(
                scope1_intensity=scope1,
                scope2_intensity=scope2,
                scope3_intensity=scope3,
                carbon_price_coverage=carbon_coverage,
                free_allocation=free_alloc,
            )

            asset = Asset(
                id=asset_id,
                name=asset_name,
                commodity=commodity,
                region=region,
                baseline_production=baseline_production,
                production_unit=production_unit,
                emissions=emissions,
                carrying_value=carrying_value,
                remaining_life_years=remaining_life,
                baseline_unit_cost=baseline_unit_cost,
                energy_cost_share=energy_cost_share,
                lat=latitude if latitude != 0.0 else None,
                lon=longitude if longitude != 0.0 else None,
            )

            companies_by_id[company_id].assets.append(asset)

        except (ValueError, TypeError) as e:
            raise ValueError(f"Assets sheet row {row_idx}: {e}")

    return list(companies_by_id.values())


# ---------------------------------------------------------------------------
# Financial-format parser  (auto-detected single-sheet financial tables)
# ---------------------------------------------------------------------------

def parse_financial_excel(path: str | Path) -> list[Company]:
    """
    Parse a simple financial-summary Excel into Company objects.

    Accepts any sheet whose first non-empty row contains columns that can be
    mapped to company name and sector.  Revenue, EBITDA, market cap, net debt,
    WACC and HQ region/country are optional — whatever is present is used.

    A synthetic "headquarters" asset is created for every company so that the
    physical-hazard engine has something to assess (GLOBAL_FALLBACK tier).
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    wb = load_workbook(path, data_only=True)
    sheet_name = _detect_financial_sheet(wb)
    if sheet_name is None:
        raise ValueError(
            "Could not find a sheet with company name and sector columns. "
            "Please use the CRI intake template (GET /template) or ensure "
            "your sheet has 'name'/'company' and 'sector'/'industry' headers."
        )

    ws = wb[sheet_name]

    # Read headers from first non-empty row
    headers: list[str] = []
    header_row_idx = 1
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        raw = [str(c).strip() for c in row if c is not None and str(c).strip()]
        if raw:
            headers = [str(c).strip() if c is not None else "" for c in row]
            header_row_idx = row_idx
            break

    if not headers:
        raise ValueError(f"Sheet '{sheet_name}' has no headers")

    # Map canonical → actual header
    col = {canonical: _find_col(headers, canonical) for canonical in _COL}

    companies: list[Company] = []
    seen_ids: set[str] = set()

    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        # Skip blank rows
        if all(v is None or str(v).strip() == "" for v in row):
            continue

        row_dict = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}

        # Company name — required
        name_col = col["company_name"]
        if name_col is None:
            continue
        raw_name = row_dict.get(name_col)
        if not raw_name or str(raw_name).strip() == "":
            continue
        company_name = str(raw_name).strip()

        # Sector — required
        sector_col = col["sector"]
        sector = str(row_dict.get(sector_col, "Other")).strip() if sector_col else "Other"
        if not sector:
            sector = "Other"

        # HQ region
        region_col = col["hq_region"]
        raw_region = str(row_dict.get(region_col, "global")).strip() if region_col else "global"
        hq_region = raw_region if raw_region else "global"

        # Financials
        rev   = _safe_float(row_dict.get(col["revenue_musd"])  if col["revenue_musd"]  else None)
        ebitda = _safe_float(row_dict.get(col["ebitda_musd"])  if col["ebitda_musd"]   else None)
        mcap   = _safe_float(row_dict.get(col["market_cap"])   if col["market_cap"]    else None) or None
        ndebt  = _safe_float(row_dict.get(col["net_debt"])     if col["net_debt"]      else None)
        wacc_v = _safe_float(row_dict.get(col["wacc_base"])    if col["wacc_base"]     else None, default=0.08)
        # Treat wacc > 1 as percentage (e.g. 8.5 → 0.085)
        if wacc_v > 1:
            wacc_v = wacc_v / 100.0

        financials = Financials(
            revenue=rev,
            ebitda=ebitda,
            capex=0.0,
            net_debt=ndebt,
            market_cap=mcap,
            wacc_base=wacc_v if wacc_v > 0 else 0.08,
        )

        # Unique ID
        base_id = _slug(company_name)
        cid = base_id
        suffix = 1
        while cid in seen_ids:
            cid = f"{base_id}-{suffix}"
            suffix += 1
        seen_ids.add(cid)

        company = Company(
            id=cid,
            name=company_name,
            sector=sector,
            hq_region=hq_region,
            financials=financials,
        )

        # Synthetic HQ asset — gives the physical-hazard engine a location
        lat_col = col["latitude"]
        lon_col = col["longitude"]
        lat = _safe_float(row_dict.get(lat_col) if lat_col else None, default=None)   # type: ignore[arg-type]
        lon = _safe_float(row_dict.get(lon_col) if lon_col else None, default=None)   # type: ignore[arg-type]

        if lat is None or lon is None:
            lat, lon = _region_centroid(hq_region)

        commodity = _sector_to_commodity(sector)
        hq_asset = Asset(
            id=f"{cid}-hq",
            name=f"{company_name} Headquarters",
            commodity=commodity,
            region=hq_region,
            baseline_production=max(rev, 1.0),  # USD-M as proxy volume
            production_unit="USD_M_revenue",
            carrying_value=mcap or (rev * 2),
            lat=lat,
            lon=lon,
            emissions=EmissionsProfile(),
        )
        company.assets.append(hq_asset)
        companies.append(company)

    if not companies:
        raise ValueError(
            f"No data rows found in sheet '{sheet_name}'. "
            "Ensure the sheet has at least one row with a company name and sector."
        )

    return companies


def auto_parse_excel(path: str | Path) -> list[Company]:
    """
    Auto-detect whether the workbook uses the CRI intake template format
    (has a 'Company' sheet) or a financial summary format (any sheet with
    name/sector columns) and dispatch accordingly.
    """
    path = Path(path)
    wb = load_workbook(path, data_only=True)
    if "Company" in wb.sheetnames:
        return parse_excel(path)
    return parse_financial_excel(path)
